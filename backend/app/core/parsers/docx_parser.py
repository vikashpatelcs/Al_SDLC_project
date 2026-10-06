"""
Parser for DOCX requirements documents.
"""
import os
from typing import List
import docx
from backend.app.core.parsers.chunk import DocumentChunk


def parse_docx_file(file_path: str, max_chunk_chars: int = 1500) -> List[DocumentChunk]:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    filename = os.path.basename(file_path)
    doc = docx.Document(file_path)

    items = []
    # Collect paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            items.append(txt)

    # Collect tables
    for t_idx, table in enumerate(doc.tables, 1):
        for row in table.rows:
            row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_txt:
                items.append(f"[Table {t_idx}] {row_txt}")

    chunks: List[DocumentChunk] = []
    current_text = []
    current_len = 0
    chunk_index = 1

    for item in items:
        if current_len + len(item) > max_chunk_chars and current_text:
            chunks.append(DocumentChunk(
                chunk_id=f"{filename}#section-{chunk_index}",
                source_file=filename,
                text="\n\n".join(current_text),
                page_or_section=f"Section {chunk_index}"
            ))
            chunk_index += 1
            current_text = [item]
            current_len = len(item)
        else:
            current_text.append(item)
            current_len += len(item)

    if current_text:
        chunks.append(DocumentChunk(
            chunk_id=f"{filename}#section-{chunk_index}",
            source_file=filename,
            text="\n\n".join(current_text),
            page_or_section=f"Section {chunk_index}"
        ))

    return chunks
